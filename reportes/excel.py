"""
reportes/excel.py — Genera el libro .xlsx de reporte de ventas (4 hojas) con openpyxl.

A diferencia del mockup web (que dibujaba letras de columna A/B/C/D a mano para simular
una hoja de cálculo), aquí se escribe un .xlsx real: Excel ya muestra sus propias
cabeceras de fila/columna, así que esa parte del mockup no se replica. Lo que sí se
replica es lo importante: fórmulas reales entre hojas, no solo valores calculados.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import DataBarRule

from reportes.data import ReporteDatos

# ── Paleta / estilos compartidos ──────────────────────────────────────────────
INK = "1E1814"
ACCENT = "8A5A3C"
HEADER_FILL = PatternFill("solid", fgColor="EDE4D6")
TOTAL_FILL = PatternFill("solid", fgColor="F0E6D3")
THIN = Side(style="thin", color="D8CBB8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOLD = Font(bold=True, color=INK)
TITLE_FONT = Font(bold=True, size=14, color=INK)
MONEY_FMT = '"$" #,##0'


def _set_widths(ws, widths: dict[str, int]):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def _header_row(ws, row: int, valores: list[str]):
    for i, val in enumerate(valores, start=1):
        c = ws.cell(row=row, column=i, value=val)
        c.font = BOLD
        c.fill = HEADER_FILL
        c.border = BORDER


def _hoja_resumen(wb: Workbook, d: ReporteDatos):
    ws = wb.active
    ws.title = "Resumen"
    _set_widths(ws, {"A": 28, "B": 18, "C": 18, "D": 14})

    ws["A1"] = f"REPORTE DE VENTAS — {d.vendedor_nombre}"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:D1")

    ws["A2"], ws["B2"] = "Vendedor", d.vendedor_nombre
    ws["C2"], ws["D2"] = "Código", d.vendedor_codigo
    ws["A3"], ws["B3"] = "Periodo", d.periodo.capitalize()
    ws["C3"], ws["D3"] = "Generado", d.generado_en.strftime("%d/%m/%Y")
    for cell in ("A2", "C2", "A3", "C3"):
        ws[cell].font = Font(color="6B5A4D")

    fila = 5
    _header_row(ws, fila, ["Indicador", "Este periodo", "Variación"])
    indicadores = [
        ("Ingresos válidos", d.ingresos_validos, d.ingresos_validos_var_pct, MONEY_FMT),
        ("Pedidos completados", d.pedidos_completados, d.pedidos_completados_var_pct, "#,##0"),
        ("Ticket promedio", d.ticket_promedio, d.ticket_promedio_var_pct, MONEY_FMT),
        ("Prendas vendidas", d.prendas_vendidas, d.prendas_vendidas_var_pct, "#,##0"),
    ]
    fila += 1
    for nombre, valor, var_pct, fmt in indicadores:
        ws.cell(row=fila, column=1, value=nombre).border = BORDER
        c = ws.cell(row=fila, column=2, value=valor)
        c.number_format = fmt
        c.border = BORDER
        c3 = ws.cell(row=fila, column=3, value=(var_pct / 100) if var_pct is not None else None)
        c3.number_format = "+0.0%;-0.0%"
        c3.border = BORDER
        if var_pct is None:
            ws.cell(row=fila, column=3, value="Sin datos del mes anterior")
        fila += 1

    fila += 1
    _header_row(ws, fila, ["Categoría", "Ingresos", "Participación"])
    fila += 1
    primera_fila_cat = fila
    for cat in d.categorias:
        ws.cell(row=fila, column=1, value=cat.nombre).border = BORDER
        c2 = ws.cell(row=fila, column=2, value=cat.ingresos)
        c2.number_format = MONEY_FMT
        c2.border = BORDER
        c3 = ws.cell(row=fila, column=3, value=cat.porcentaje / 100)
        c3.number_format = "0.0%"
        c3.border = BORDER
        fila += 1
    ultima_fila_cat = fila - 1

    if d.categorias:
        # Barra de datos nativa de Excel sobre la columna de participación.
        rango = f"C{primera_fila_cat}:C{ultima_fila_cat}"
        ws.conditional_formatting.add(
            rango,
            DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=ACCENT),
        )
        ws.cell(row=fila, column=1, value="Total").font = BOLD
        ws.cell(row=fila, column=1).fill = TOTAL_FILL
        c = ws.cell(row=fila, column=2, value=f"=SUM(B{primera_fila_cat}:B{ultima_fila_cat})")
        c.number_format = MONEY_FMT
        c.font = BOLD
        c.fill = TOTAL_FILL
        ws.cell(row=fila, column=3, value=1).number_format = "0.0%"
        ws.cell(row=fila, column=3).fill = TOTAL_FILL
    else:
        ws.cell(row=fila, column=1, value="Sin ventas completadas en este periodo.").font = Font(italic=True, color="6B5A4D")


def _hoja_transacciones(wb: Workbook, d: ReporteDatos, comision_pct: float):
    ws = wb.create_sheet("Transacciones")
    _set_widths(ws, {"A": 12, "B": 10, "C": 30, "D": 16, "E": 18, "F": 13, "G": 13, "H": 13})
    _header_row(ws, 1, ["Fecha", "Pedido", "Producto", "Categoría", "Comprador", "Precio", "Estado", "Neto"])

    fila = 2
    for t in d.transacciones:
        ws.cell(row=fila, column=1, value=t.fecha.strftime("%d/%m/%Y")).border = BORDER
        ws.cell(row=fila, column=2, value=t.id_pedido).border = BORDER
        ws.cell(row=fila, column=3, value=t.producto).border = BORDER
        ws.cell(row=fila, column=4, value=t.categoria).border = BORDER
        ws.cell(row=fila, column=5, value=t.comprador).border = BORDER
        c_precio = ws.cell(row=fila, column=6, value=t.precio)
        c_precio.number_format = MONEY_FMT
        c_precio.border = BORDER
        ws.cell(row=fila, column=7, value=t.estado).border = BORDER
        c_neto = ws.cell(
            row=fila, column=8,
            value=f'=IF(G{fila}="completado",F{fila}*(1-{comision_pct/100}),IF(G{fila}="cancelado",0,""))',
        )
        c_neto.number_format = MONEY_FMT
        c_neto.border = BORDER
        c_neto.fill = TOTAL_FILL
        fila += 1

    if d.transacciones:
        ws.cell(row=fila, column=1, value=f"Total ({len(d.transacciones)} pedidos)").font = BOLD
        c = ws.cell(row=fila, column=6, value=f"=SUM(F2:F{fila-1})")
        c.number_format = MONEY_FMT
        c.font = BOLD
        c2 = ws.cell(row=fila, column=8, value=f"=SUM(H2:H{fila-1})")
        c2.number_format = MONEY_FMT
        c2.font = BOLD
    else:
        ws.cell(row=fila, column=1, value="Sin transacciones en este periodo.").font = Font(italic=True, color="6B5A4D")

    ws.cell(row=1, column=1).alignment = Alignment(wrap_text=False)
    ws.freeze_panes = "A2"


def _hoja_prendas(wb: Workbook, d: ReporteDatos):
    ws = wb.create_sheet("Prendas")
    _set_widths(ws, {"A": 14, "B": 30, "C": 16, "D": 14, "E": 8, "F": 14, "G": 14, "H": 12, "I": 13, "J": 13})
    _header_row(ws, 1, ["SKU", "Prenda", "Categoría", "Marca", "Talla", "Color", "Condición", "Pedido", "Precio", "Estado"])

    fila = 2
    for t in d.transacciones:
        sku = f"VT-{t.id_prenda}" if t.id_prenda is not None else "—"
        ws.cell(row=fila, column=1, value=sku).border = BORDER
        ws.cell(row=fila, column=2, value=t.producto).border = BORDER
        ws.cell(row=fila, column=3, value=t.categoria).border = BORDER
        ws.cell(row=fila, column=4, value=t.marca).border = BORDER
        ws.cell(row=fila, column=5, value=t.talla).border = BORDER
        ws.cell(row=fila, column=6, value=t.color).border = BORDER
        ws.cell(row=fila, column=7, value=t.condicion).border = BORDER
        ws.cell(row=fila, column=8, value=t.id_pedido).border = BORDER
        c = ws.cell(row=fila, column=9, value=t.precio)
        c.number_format = MONEY_FMT
        c.border = BORDER
        ws.cell(row=fila, column=10, value=t.estado).border = BORDER
        fila += 1

    if d.transacciones:
        ws.cell(row=fila, column=1, value=f"{len(d.transacciones)} prendas en el periodo").font = BOLD
        c = ws.cell(row=fila, column=9, value=f"=SUM(I2:I{fila-1})")
        c.number_format = MONEY_FMT
        c.font = BOLD
    else:
        ws.cell(row=fila, column=1, value="Sin prendas vendidas en este periodo.").font = Font(italic=True, color="6B5A4D")

    ws.freeze_panes = "A2"


def _hoja_liquidacion(wb: Workbook, d: ReporteDatos):
    ws = wb.create_sheet("Liquidación")
    _set_widths(ws, {"A": 30, "B": 18, "C": 40})

    ws["A1"] = f"LIQUIDACIÓN {d.periodo.upper()} — {d.vendedor_codigo}"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:C1")

    _header_row(ws, 2, ["Concepto", "Valor", "Nota"])

    # Columna B lleva la FÓRMULA REAL (Excel la recalcula si cambia Transacciones);
    # la columna C es solo una nota en lenguaje llano, no un duplicado del cálculo.
    filas = [
        ("Ventas brutas (completadas + canceladas)",
         '=SUMIF(Transacciones!G:G,"completado",Transacciones!F:F)+SUMIF(Transacciones!G:G,"cancelado",Transacciones!F:F)',
         "Suma de pedidos completados y cancelados"),
        ("Ventas canceladas",
         '=-SUMIF(Transacciones!G:G,"cancelado",Transacciones!F:F)',
         "Se resta de las ventas brutas"),
        ("Ventas válidas", "=B3+B4", "Brutas menos canceladas"),
        (f"Comisión VINT ({d.comision_pct:.0f}%)", f"=-B5*{d.comision_pct/100}", "Sobre las ventas válidas"),
    ]
    fila = 3
    for nombre, formula, nota in filas:
        ws.cell(row=fila, column=1, value=nombre).border = BORDER
        c = ws.cell(row=fila, column=2, value=formula)
        c.number_format = MONEY_FMT
        c.border = BORDER
        nc = ws.cell(row=fila, column=3, value=nota)
        nc.font = Font(italic=True, color="6B5A4D", size=10)
        nc.border = BORDER
        fila += 1

    ws.cell(row=fila, column=1, value="Total neto a recibir").font = Font(bold=True, size=13, color=ACCENT)
    ws.cell(row=fila, column=1).fill = TOTAL_FILL
    c = ws.cell(row=fila, column=2, value="=B5+B6")
    c.number_format = MONEY_FMT
    c.font = Font(bold=True, size=13, color=ACCENT)
    c.fill = TOTAL_FILL
    nc = ws.cell(row=fila, column=3, value="Ventas válidas menos comisión")
    nc.font = Font(italic=True, color="6B5A4D", size=10)
    nc.fill = TOTAL_FILL
    fila += 2

    if d.pedidos_en_curso > 0:
        ws.cell(
            row=fila, column=1,
            value=f"{d.pedidos_en_curso} pedido(s) en curso (pendiente/enviado) no incluidos en esta liquidación.",
        ).font = Font(italic=True, color="6B5A4D", size=10)

    ws.cell(row=fila + 1, column=1,
            value="Toda la hoja depende de Transacciones: si cambia el estado de un pedido, la liquidación se recalcula sola.").font = Font(italic=True, color="6B5A4D", size=10)


def generar_excel(d: ReporteDatos) -> BytesIO:
    from config import get_settings
    comision_pct = get_settings().VINT_COMISION_PORCENTAJE * 100

    wb = Workbook()
    _hoja_resumen(wb, d)
    _hoja_transacciones(wb, d, comision_pct)
    _hoja_prendas(wb, d)
    _hoja_liquidacion(wb, d)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
